import base64
import re
import zlib
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from school.constants import Day
from school.tests.factories import make_user
from scheduling.tests.fixtures import SchoolFixture

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def pdf_text(content):
    """All text drawn in a PDF. reportlab encodes page streams as ASCII85 + Flate."""
    chunks = []
    for stream in re.findall(rb"stream\r?\n(.*?)\r?\n?endstream", content, re.S):
        data = stream.strip()
        try:
            if data.endswith(b"~>"):
                data = base64.a85decode(data[:-2])
            chunks.append(zlib.decompress(data))
        except (ValueError, zlib.error):
            chunks.append(stream)
    return b"".join(chunks).decode("latin-1")


class ExportTests(SchoolFixture):
    def setUp(self):
        self.assign(self.c10a, Day.MONDAY, 1)  # Ms Maths, Mathematics
        self.assign(self.c10b, Day.TUESDAY, 3, teacher=self.mr_physics)
        self.client.force_login(self.viewer)

    def test_excel_workbook(self):
        response = self.client.get(reverse("exports:excel"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX)
        self.assertIn("attachment;", response["Content-Disposition"])
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["10A", "10B", "Teachers", "Workload"])
        values = [c.value for row in wb["10A"].iter_rows() for c in row if c.value]
        self.assertIn("MATH\nMs Maths", values)
        self.assertIn("Monday", values)
        teacher_values = [c.value for row in wb["Teachers"].iter_rows() for c in row if c.value]
        self.assertIn("10B\nMATH", teacher_values)
        workload = {row[0].value: row for row in wb["Workload"].iter_rows(min_row=4)}
        self.assertEqual(workload["Ms Maths"][2].value, "1/3")  # Monday
        self.assertEqual(workload["Mr Physics"][3].value, "1/5")  # Tuesday

    def test_class_pdf(self):
        response = self.client.get(reverse("exports:pdf_class", args=[self.c10a.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        text = pdf_text(response.content)
        self.assertIn("Class 10A", text)
        self.assertIn("Ms Maths", text)
        self.assertNotIn("Mr Physics", text)

    def test_teacher_pdf(self):
        response = self.client.get(reverse("exports:pdf_teacher", args=[self.mr_physics.pk]))
        self.assertEqual(response["Content-Type"], "application/pdf")
        text = pdf_text(response.content)
        self.assertIn("Mr Physics", text)
        self.assertIn("10B", text)

    def test_school_pdf(self):
        response = self.client.get(reverse("exports:pdf_school"))
        self.assertEqual(response["Content-Type"], "application/pdf")
        text = pdf_text(response.content)
        for expected in ("Class 10A", "Class 10B", "Ms Maths", "Mr Physics", "Teacher workloads"):
            self.assertIn(expected, text)

    def test_print_pages(self):
        response = self.client.get(reverse("exports:print_classes"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/html; charset=utf-8")
        self.assertContains(response, 'class="print-page"', count=2)
        self.assertContains(response, "Ms Maths")
        self.assertContains(response, "css/print.css")
        response = self.client.get(reverse("exports:print_class", args=[self.c10a.pk]))
        self.assertContains(response, 'class="print-page"', count=1)
        self.assertContains(response, reverse("scheduling:class_timetable", args=[self.c10a.pk]))  # back link
        self.assertNotContains(response, '<button type="button" class="slot')
        response = self.client.get(reverse("exports:print_teacher", args=[self.mr_physics.pk]))
        self.assertContains(response, "10B")

    def test_missing_objects_are_404(self):
        for name in ("exports:pdf_class", "exports:pdf_teacher", "exports:print_class", "exports:print_teacher"):
            with self.subTest(name=name):
                self.assertEqual(self.client.get(reverse(name, args=[999])).status_code, 404)

    def test_exports_respect_permissions(self):
        urls = [
            reverse("exports:excel"),
            reverse("exports:pdf_school"),
            reverse("exports:pdf_class", args=[self.c10a.pk]),
            reverse("exports:pdf_teacher", args=[self.ms_maths.pk]),
            reverse("exports:print_classes"),
        ]
        self.client.logout()
        for url in urls:
            with self.subTest(anonymous=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 302)
                self.assertIn(reverse("login"), response["Location"])
        self.client.force_login(make_user("nobody"))
        for url in urls:
            with self.subTest(no_role=url):
                self.assertEqual(self.client.get(url).status_code, 403)
        for user in (self.viewer, self.editor):
            self.client.force_login(user)
            for url in urls:
                with self.subTest(user=user.username, url=url):
                    self.assertEqual(self.client.get(url).status_code, 200)

    def test_empty_school_exports(self):
        from scheduling import services

        services.reset_class(class_group=self.c10a, user=None)
        services.reset_class(class_group=self.c10b, user=None)
        self.assertEqual(self.client.get(reverse("exports:excel")).status_code, 200)
        self.assertEqual(self.client.get(reverse("exports:pdf_school")).status_code, 200)
