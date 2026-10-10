"""
Reading a teacher allotment workbook (e.g. "TT Allotment 25-26.xlsx").

Layout of the "Allotment" sheet:

    Teacher            | Class        | Subject | Periods/week
    A B Bhatt          |              |         | 22            <- teacher row (total)
                       | VII A        | HIS     | 3             <- one allocation
                       | VII C,VII D  | PHE     | 2             <- several classes at once

`parse_allotment()` returns every allocation. `plan_import()` splits them into
what the system can represent today (one teacher, one class, one subject per
slot: CLAUDE.md rule 4) and what it cannot yet:

- combined lessons: one row covering several classes in the same slot;
- parallel groups: several teachers taking the same class and subject at once
  (labs, language groups, split activity classes).

Skipped allocations are returned with the reason so they can be reported.
"""

from collections import defaultdict
from dataclasses import dataclass, field

from openpyxl import load_workbook

SHEET = "Allotment"

# Grade order as printed on the sheet: Sr Nur, Prep, I … X.
GRADE_ORDER = ["Sr Nur", "Prep", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]

# Subject names for codes whose meaning is clear from the workbook. Unknown
# codes keep the code as their name; rename them later in the admin.
SUBJECT_NAMES = {
    "A&C": "Art & Craft",
    "AI": "Artificial Intelligence",
    "BIO": "Biology",
    "CHE": "Chemistry",
    "CIV": "Civics",
    "COM": "Computer Science",
    "ECO": "Economics",
    "ENG": "English",
    "EVS": "Environmental Studies",
    "GEO": "Geography",
    "HC": "Happiness Class",
    "HIN": "Hindi",
    "HIN/URD": "Hindi / Urdu",
    "HIS": "History",
    "LIB": "Library",
    "MAT": "Mathematics",
    "MUS": "Music",
    "PHE": "Physical Education",
    "PHY": "Physics",
    "ROB": "Robotics",
    "SCI": "Science",
    "ScL": "Science Lab",
    "SKT": "Skating",
    "URD": "Urdu",
}

COMBINED = "combined"
PARALLEL = "parallel"
REASONS = {
    COMBINED: "one lesson for several classes at once (combined lessons are not supported yet)",
    PARALLEL: "several teachers take this class and subject at the same time (parallel groups are not supported yet)",
}


class AllotmentError(Exception):
    pass


@dataclass(frozen=True)
class Allocation:
    row: int
    teacher: str
    classes: tuple
    subject: str
    periods: int


@dataclass
class ImportPlan:
    teachers: list  # names, in sheet order
    classes: list  # labels, in grade order
    subjects: dict  # code -> name
    qualifications: set  # {(teacher, subject code)}
    requirements: dict  # (class label, subject code) -> (periods, teacher)
    skipped: list = field(default_factory=list)  # [(Allocation, reason code)]


def split_class(label):
    """'VII A' -> ('VII', 'A'); 'Sr Nur B' -> ('Sr Nur', 'B')."""
    grade, _, section = label.strip().rpartition(" ")
    if not grade or not section:
        raise AllotmentError(f"Cannot read the class “{label}”: expected a grade and a section, e.g. “VII A”.")
    return grade, section


def grade_order(grade):
    return GRADE_ORDER.index(grade) if grade in GRADE_ORDER else len(GRADE_ORDER)


def parse_allotment(path):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable file
        raise AllotmentError(f"Cannot open the workbook: {exc}") from exc
    if SHEET not in workbook.sheetnames:
        raise AllotmentError(f"The workbook has no “{SHEET}” sheet (found: {', '.join(workbook.sheetnames)}).")

    allocations, errors = [], []
    teacher = None
    for number, row in enumerate(workbook[SHEET].iter_rows(min_row=2, max_col=4, values_only=True), start=2):
        name, classes, subject, periods = (str(v).strip() if v is not None else "" for v in row)
        if not any((name, classes, subject, periods)):
            continue
        if name:
            teacher = " ".join(name.split())
            continue
        if not (classes and subject and periods):
            errors.append(f"Row {number}: class, subject and periods are all needed.")
            continue
        if teacher is None:
            errors.append(f"Row {number}: an allocation appears before any teacher name.")
            continue
        try:
            count = int(float(periods))
        except ValueError:
            errors.append(f"Row {number}: periods “{periods}” is not a number.")
            continue
        labels = tuple(" ".join(part.split()) for part in classes.split(",") if part.strip())
        for label in labels:
            try:
                split_class(label)
            except AllotmentError as exc:
                errors.append(f"Row {number}: {exc}")
        allocations.append(Allocation(number, teacher, labels, subject, count))
    workbook.close()
    if errors:
        raise AllotmentError("The workbook has problems:\n" + "\n".join(errors))
    return allocations


def plan_import(allocations):
    teachers, classes, subjects = [], set(), {}
    qualifications = set()
    by_class_subject = defaultdict(list)
    for a in allocations:
        if a.teacher not in teachers:
            teachers.append(a.teacher)
        classes.update(a.classes)
        subjects.setdefault(a.subject, SUBJECT_NAMES.get(a.subject, a.subject))
        qualifications.add((a.teacher, a.subject))
        for label in a.classes:
            by_class_subject[(label, a.subject)].append(a)

    requirements, skipped = {}, []
    seen_skipped = set()
    for key, group in by_class_subject.items():
        if any(len(a.classes) > 1 for a in group):
            # Every row for this class and subject is part of the combined lesson.
            for a in group:
                if a not in seen_skipped:
                    seen_skipped.add(a)
                    skipped.append((a, COMBINED))
            continue
        if len(group) > 1:
            for a in group:
                if a not in seen_skipped:
                    seen_skipped.add(a)
                    skipped.append((a, PARALLEL))
            continue
        a = group[0]
        requirements[key] = (a.periods, a.teacher)

    def class_key(label):
        grade, section = split_class(label)
        return (grade_order(grade), grade, section)

    skipped.sort(key=lambda item: item[0].row)
    return ImportPlan(
        teachers=teachers,
        classes=sorted(classes, key=class_key),
        subjects=subjects,
        qualifications=qualifications,
        requirements=requirements,
        skipped=skipped,
    )
