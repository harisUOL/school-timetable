"""Excel workbook: one sheet per class, a Teachers sheet and a Workload sheet (openpyxl)."""

import re
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import tables

HEADER_FILL = PatternFill("solid", fgColor="DCE6F5")
RECESS_FILL = PatternFill("solid", fgColor="EEEEEE")
THIN = Side(style="thin", color="B7C0CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _sheet_title(name, used):
    title = re.sub(r"[\[\]:*?/\\]", "-", name)[:31] or "Sheet"
    base, n = title, 2
    while title in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def _write_table(ws, table, start_row=1, cell_height=32):
    ws.cell(start_row, 1, table.title).font = Font(bold=True, size=13)
    ws.cell(start_row + 1, 1, table.subtitle).font = Font(italic=True, color="555555")
    header_row = start_row + 2
    for col, label in enumerate(table.header, start=1):
        cell = ws.cell(header_row, col, label)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        cell.border = BORDER
    for r, row in enumerate(table.rows, start=header_row + 1):
        ws.row_dimensions[r].height = cell_height
        for col, value in enumerate(row, start=1):
            cell = ws.cell(r, col, value)
            cell.alignment = WRAP if col > 1 else Alignment(vertical="center")
            cell.border = BORDER
            if col == 1:
                cell.font = Font(bold=True)
            if (col - 1) in table.recess_columns:
                cell.fill = RECESS_FILL
    return header_row + len(table.rows) + 2


def _widths(ws, first=14, rest=14):
    ws.column_dimensions["A"].width = first
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = rest


def build_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for class_group in tables.all_classes():
        ws = wb.create_sheet(_sheet_title(str(class_group), used))
        _write_table(ws, tables.class_table(class_group))
        _widths(ws)
        ws.freeze_panes = "B4"

    ws = wb.create_sheet(_sheet_title("Teachers", used))
    row = 1
    for teacher in tables.active_teachers():
        row = _write_table(ws, tables.teacher_table(teacher), start_row=row)
    _widths(ws, first=14, rest=13)

    ws = wb.create_sheet(_sheet_title("Workload", used))
    _write_table(ws, tables.workload_table(), cell_height=18)
    _widths(ws, first=22, rest=9)
    ws.column_dimensions["B"].width = 34
    ws.freeze_panes = "C4"

    wb.properties.title = f"{tables.school_name()} timetable"
    wb.properties.created = timezone.now().replace(tzinfo=None)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
